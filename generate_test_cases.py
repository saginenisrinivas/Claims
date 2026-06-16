import openpyxl
from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side)
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── Colour palette ────────────────────────────────────────────────────────────
HDR_FILL   = PatternFill("solid", fgColor="0F3460")   # dark navy
MOD_FILL   = PatternFill("solid", fgColor="1A5276")   # module row – deep blue
ALT_FILL   = PatternFill("solid", fgColor="EEF2FF")   # alternating light row
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
PASS_FILL  = PatternFill("solid", fgColor="D4EDDA")
FAIL_FILL  = PatternFill("solid", fgColor="F8D7DA")
BLOK_FILL  = PatternFill("solid", fgColor="FFF3CD")

HDR_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
MOD_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", bold=True, size=10)

THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ─────────────────────────────────────────────────────────────────────────────
# Sheet 1 – Test Cases
# ─────────────────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "Test Cases"

COLS = ["TC #", "Module", "Test Case Description", "Pre-conditions",
        "Test Steps", "Expected Result", "Test Data", "Status", "Remarks"]
COL_W = [7, 22, 32, 22, 45, 45, 28, 10, 18]

# header row
for ci, (label, w) in enumerate(zip(COLS, COL_W), 1):
    cell = ws.cell(row=1, column=ci, value=label)
    cell.fill   = HDR_FILL
    cell.font   = HDR_FONT
    cell.border = BORDER
    cell.alignment = CENTER
    ws.column_dimensions[get_column_letter(ci)].width = w

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

# ── Test data ────────────────────────────────────────────────────────────────
# Each entry: (tc, module, description, preconditions, steps, expected, data)
TESTS = [
    # ── MODULE 1: Login ──────────────────────────────────────────────────────
    ("MOD", "MODULE 1: Login & Authentication", "", "", "", "", ""),
    ("TC-01","Login","Valid login – Claims Officer",
     "App running at localhost:9090",
     "1. Open http://localhost:9090\n2. Enter officer@claims.com / officer123\n3. Click Login",
     "Redirected to Dashboard. Sidebar shows: Dashboard, Case Registration, Case Acceptance, Case Evaluation.",
     "Email: officer@claims.com\nPwd: officer123"),
    ("TC-02","Login","Valid login – Case Manager",
     "App running",
     "1. Enter manager@claims.com / manager123\n2. Click Login",
     "Redirected to Dashboard. Sidebar shows: Dashboard, Case Approval only.",
     "Email: manager@claims.com\nPwd: manager123"),
    ("TC-03","Login","Valid login – Administrator",
     "App running",
     "1. Enter admin@claims.com / admin123\n2. Click Login",
     "Redirected to Dashboard.",
     "Email: admin@claims.com\nPwd: admin123"),
    ("TC-04","Login","Invalid credentials",
     "App running",
     "1. Enter any email with wrong password\n2. Click Login",
     "Error message displayed. User stays on login page.",
     "Email: officer@claims.com\nPwd: wrongpass"),
    ("TC-05","Login","Blank fields submission",
     "App running",
     "1. Leave email and password blank\n2. Click Login",
     "Validation error shown on screen.",
     "Leave all fields empty"),
    ("TC-06","Login","Session expiry",
     "Logged in",
     "1. Log in\n2. Leave browser idle for 30+ minutes\n3. Navigate to any page",
     "Redirected back to login page.",
     "Wait 30 min"),
    ("TC-07","Login","Sign Out",
     "Logged in as any user",
     "1. Click 'Sign Out' in sidebar footer",
     "Redirected to login page. Session cleared.",
     "—"),

    # ── MODULE 2: Dashboard ──────────────────────────────────────────────────
    ("MOD", "MODULE 2: Dashboard", "", "", "", "", ""),
    ("TC-08","Dashboard","Stats cards load correctly",
     "Logged in as Claims Officer",
     "1. Open Dashboard",
     "6 stat cards visible: Total Cases, Pending Acceptance, Under Evaluation, Awaiting Approval, Approved, Rejected/Denied with correct counts.",
     "—"),
    ("TC-09","Dashboard","Claim Type shows label in Recent Cases",
     "At least one case registered with Funeral or Death type",
     "1. Open Dashboard\n2. Check Type column in Recent Cases table",
     "Shows 'Funeral' or 'Death', NOT raw codes '15' or '2'.",
     "—"),
    ("TC-10","Dashboard","Pipeline chart displays",
     "Logged in",
     "1. Open Dashboard\n2. View Pipeline section",
     "Horizontal bar chart shows count for each stage: Submitted, Accepted, Under Evaluation, Evaluated, Approved.",
     "—"),
    ("TC-11","Dashboard","New Case button – Claims Officer",
     "Logged in as Claims Officer",
     "1. Open Dashboard\n2. Check top-right header",
     "'New Case' button is visible.",
     "—"),
    ("TC-12","Dashboard","New Case button – Case Manager",
     "Logged in as Case Manager",
     "1. Open Dashboard\n2. Check top-right header",
     "'New Case' button is NOT visible.",
     "—"),
    ("TC-13","Dashboard","Quick Actions by role",
     "Logged in as Claims Officer",
     "1. Open Dashboard\n2. Check Quick Actions panel",
     "Shows: Register New Case, Review Pending Cases, Cases for Evaluation.",
     "—"),

    # ── MODULE 3: Case Registration ──────────────────────────────────────────
    ("MOD", "MODULE 3: Case Registration – Validations", "", "", "", "", ""),
    ("TC-14","Case Registration","Submit with all fields empty",
     "Logged in as Claims Officer",
     "1. Open Case Registration\n2. Click Submit without filling any fields",
     "Red error messages appear under all required fields. Alert banner: 'Please fill all required fields.'",
     "Leave all fields blank"),
    ("TC-15","Case Registration","Life Assured Name required",
     "Logged in as Claims Officer",
     "1. Fill all required fields except Life Assured Name\n2. Click Submit",
     "Error shown under Life Assured Name field only.",
     "Leave Life Assured Name blank"),
    ("TC-16","Case Registration","Event Date required",
     "Logged in as Claims Officer",
     "1. Fill all fields except Event Date\n2. Click Submit",
     "Red error 'Required.' shown under Event Date.",
     "Leave Event Date blank"),
    ("TC-17","Case Registration","Email format validation",
     "Logged in as Claims Officer",
     "1. Enter 'notanemail' in Email Address\n2. Click Submit",
     "Error: 'Valid email required.'",
     "Email: notanemail"),
    ("TC-18","Case Registration","Future Event Date blocked",
     "Logged in as Claims Officer",
     "1. Fill all fields, set Event Date to tomorrow\n2. Click Submit",
     "Server returns error: 'Event date cannot be in the future.'",
     "Event Date: tomorrow's date"),

    ("MOD", "MODULE 3: Case Registration – Dropdowns", "", "", "", "", ""),
    ("TC-19","Case Registration","ID Type dropdown values",
     "Logged in as Claims Officer, Case Registration open",
     "1. Click ID Type dropdown\n2. Check all options",
     "Shows exactly 3 options: Omang, Others, Passport. No National ID. No Driver's Licence.",
     "—"),
    ("TC-20","Case Registration","Gender dropdown values",
     "Logged in as Claims Officer",
     "1. Click Gender dropdown\n2. Check all options",
     "Shows exactly 2 options: Male, Female.",
     "—"),
    ("TC-21","Case Registration","Type of Claim dropdown values",
     "Logged in as Claims Officer",
     "1. Click Type of Claim dropdown",
     "Shows: Funeral, Death only.",
     "—"),
    ("TC-22","Case Registration","Claim Nature dropdown values",
     "Logged in as Claims Officer",
     "1. Click Claim Nature dropdown",
     "Shows: Illness (Non-accident), Accident, Suicide.",
     "—"),
    ("TC-23","Case Registration","Relation dropdown values",
     "Logged in as Claims Officer",
     "1. Click Relation With Life Assured dropdown\n2. Count and verify options",
     "Shows 19 values: N/A, Spouse, Child, Parent, Relative, Self, Other, Employer, Employee, Heir, Trustee, Official Assignee, Public Trustee, Siblings, Niece/Nephew, Assignee, Administrator, Executrix, Executor.",
     "—"),

    ("MOD", "MODULE 3: Case Registration – Policy Lookup", "", "", "", "", ""),
    ("TC-24","Case Registration","Fetch Policies without ID No.",
     "Logged in as Claims Officer",
     "1. Leave ID No. blank\n2. Click 'Fetch Policies'",
     "Toast warning: 'Please enter an ID No. in Insured Information first.'",
     "ID No.: blank"),
    ("TC-25","Case Registration","Fetch Policies with valid ID",
     "Logged in as Claims Officer, valid ID available",
     "1. Enter valid ID No.\n2. Click 'Fetch Policies'",
     "Policy table appears. Active policies show 'Select' button. Inactive policies show 'Not eligible'.",
     "ID No.: valid insured ID"),
    ("TC-26","Case Registration","Select Active policy",
     "Policy results loaded",
     "1. Click 'Select' on an Active/Inforce policy",
     "Row highlighted. Banner shows selected policy. Life Assured Name pre-fills from policy holder name.",
     "—"),
    ("TC-27","Case Registration","Submit without selecting policy",
     "ID No. entered but no policy selected",
     "1. Enter ID No.\n2. Do NOT select a policy\n3. Fill other required fields\n4. Click Submit",
     "Error: 'Please fetch and select a policy before submitting.'",
     "—"),
    ("TC-28","Case Registration","Submit with non-Active policy",
     "Policy with Lapsed/Cancelled status available",
     "1. Fetch policies\n2. (If possible) select Lapsed policy\n3. Submit",
     "Error: 'Registration is not allowed. Policy status is not Active.'",
     "—"),

    ("MOD", "MODULE 3: Case Registration – Documents", "", "", "", "", ""),
    ("TC-29","Case Registration","Browse without selecting document type",
     "Logged in as Claims Officer",
     "1. Click 'Browse & Add' without selecting doc type from dropdown",
     "Dropdown border turns red. Toast: 'Please select a document type before choosing a file.'",
     "—"),
    ("TC-30","Case Registration","Upload valid PDF",
     "Doc type selected",
     "1. Select 'Death Certificate' from dropdown\n2. Click 'Browse & Add'\n3. Select a PDF file under 10 MB",
     "File appears in staged list with type badge and 'Pending upload' status.",
     "Any PDF < 10MB"),
    ("TC-31","Case Registration","Upload file exceeding 10 MB",
     "Doc type selected",
     "1. Select doc type\n2. Click Browse\n3. Select file > 10MB",
     "Toast error: 'File exceeds 10 MB limit.' File not added.",
     "File > 10MB"),
    ("TC-32","Case Registration","Remove staged document",
     "At least one file staged",
     "1. Click ✕ next to a staged file",
     "File removed from staged list immediately.",
     "—"),

    ("MOD", "MODULE 3: Case Registration – THITO Integration", "", "", "", "", ""),
    ("TC-33","Case Registration","Successful submission – loading spinner",
     "All required fields filled, THITO API reachable",
     "1. Fill all required fields\n2. Click Submit\n3. Observe Submit button immediately",
     "Submit button shows spinning icon and 'Submitting...' text. Button is disabled.",
     "Claim Type: Funeral\nEvent Date: today or past"),
    ("TC-34","Case Registration","Success modal shown",
     "Submission succeeds",
     "1. Complete a valid submission",
     "Success modal appears with: System Case No. (CLM-YYYY-XXXX), THITO Case Number (e.g. FUN260002276), THITO Status (CASE REGISTERED and WAITING FOR ACCEPTANCE).",
     "—"),
    ("TC-35","Case Registration","THITO Case Number prefix for Funeral",
     "Funeral claim submitted",
     "1. Submit a Funeral claim\n2. Check THITO Case Number in modal",
     "THITO Case Number starts with 'FUN'.",
     "Claim Type: Funeral"),
    ("TC-36","Case Registration","THITO Case Number prefix for Death",
     "Death claim submitted",
     "1. Submit a Death claim\n2. Check THITO Case Number in modal",
     "THITO Case Number starts with 'DEA'.",
     "Claim Type: Death"),
    ("TC-37","Case Registration","Continue with Acceptance button",
     "Success modal open",
     "1. After successful submission\n2. Click 'Continue with Acceptance'",
     "Navigates to Case Acceptance page. Detail modal automatically opens for the newly registered case.",
     "—"),
    ("TC-38","Case Registration","Close modal resets form",
     "Success modal open",
     "1. Click 'Close' in success modal",
     "Modal closes. All form fields reset to blank/defaults. Ready for new entry.",
     "—"),
    ("TC-39","Case Registration","Submit with staged documents",
     "Documents staged before submit",
     "1. Stage 2 documents\n2. Fill required fields\n3. Submit",
     "After registration succeeds, documents upload one by one showing 'Uploading...' then '✓ Uploaded'.",
     "2 PDF files < 10MB each"),
    ("TC-40","Case Registration","THITO API failure shows error",
     "THITO API unavailable or rejecting data",
     "1. Submit with invalid/missing SOAP fields\n2. Observe result",
     "Error message from THITO API displayed on screen. Case NOT saved in local database.",
     "—"),

    # ── MODULE 4: Case Acceptance ─────────────────────────────────────────────
    ("MOD", "MODULE 4: Case Acceptance", "", "", "", "", ""),
    ("TC-41","Case Acceptance","THITO Case Number column visible",
     "Logged in as Claims Officer, at least one registered case",
     "1. Open Case Acceptance page\n2. Observe table columns",
     "Column 'THITO Case Number' appears between Case ID and Life Assured. Shows values like 'FUN260002276'. Seeded cases show '—'.",
     "—"),
    ("TC-42","Case Acceptance","Claim Type shows label",
     "Cases with Funeral or Death type exist",
     "1. Open Case Acceptance\n2. Check Type column",
     "Shows 'Funeral' or 'Death', NOT '15' or '2'.",
     "—"),
    ("TC-43","Case Acceptance","Search by name",
     "Cases loaded",
     "1. Type a name in the search box",
     "List filters in real time to matching names.",
     "Search: any known Life Assured name"),
    ("TC-44","Case Acceptance","Filter by Submitted status",
     "Cases in multiple statuses exist",
     "1. Select 'Submitted' from Status filter",
     "Only SUBMITTED cases displayed.",
     "—"),
    ("TC-45","Case Acceptance","View case detail modal",
     "At least one case exists",
     "1. Click 'View' on any case",
     "Modal opens showing: THITO Case Number, THITO Status, Claim Type label (e.g. Funeral), Claim Nature label (e.g. Accident).",
     "—"),
    ("TC-46","Case Acceptance","Accept a Submitted case",
     "A SUBMITTED case exists",
     "1. Click 'Accept' on a SUBMITTED case\n2. Optionally enter notes\n3. Click 'Confirm Acceptance'",
     "Toast: 'Case accepted successfully.' Case status changes to ACCEPTED. Accept/Reject buttons disappear.",
     "—"),
    ("TC-47","Case Acceptance","Reject without notes",
     "A SUBMITTED case exists",
     "1. Click 'Reject'\n2. Leave notes blank\n3. Click 'Confirm Rejection'",
     "Error: 'Rejection reason is required.' Case not rejected.",
     "Notes: blank"),
    ("TC-48","Case Acceptance","Reject with notes",
     "A SUBMITTED case exists",
     "1. Click 'Reject'\n2. Enter rejection reason\n3. Click 'Confirm Rejection'",
     "Toast: 'Case rejected.' Status changes to REJECTED.",
     "Notes: 'Pre-existing condition not covered.'"),
    ("TC-49","Case Acceptance","No action buttons for non-Submitted",
     "An ACCEPTED or REJECTED case exists",
     "1. View an ACCEPTED case\n2. Check footer of modal",
     "Accept and Reject buttons NOT visible.",
     "—"),
    ("TC-50","Case Acceptance","Auto-open via Continue with Acceptance",
     "Just completed Case Registration",
     "1. Click 'Continue with Acceptance' in registration success modal",
     "Case Acceptance page loads. Detail modal opens automatically for the new case.",
     "—"),

    # ── MODULE 5: Case Evaluation ─────────────────────────────────────────────
    ("MOD", "MODULE 5: Case Evaluation", "", "", "", "", ""),
    ("TC-51","Case Evaluation","Only relevant cases shown",
     "Cases in various statuses exist",
     "1. Open Case Evaluation",
     "Only ACCEPTED, UNDER_EVALUATION, EVALUATED cases shown. SUBMITTED/REJECTED not visible.",
     "—"),
    ("TC-52","Case Evaluation","Start Evaluation",
     "An ACCEPTED case exists",
     "1. Click 'Start Eval' on ACCEPTED case",
     "Toast: 'Evaluation started.' Status changes to UNDER_EVALUATION.",
     "—"),
    ("TC-53","Case Evaluation","Evaluate without fields",
     "A case is UNDER_EVALUATION",
     "1. Click 'Evaluate'\n2. Leave amount, risk, notes blank\n3. Click Submit",
     "Red error messages shown under Amount, Risk Level, Notes.",
     "—"),
    ("TC-54","Case Evaluation","Notes under 20 characters",
     "Evaluation form open",
     "1. Enter fewer than 20 characters in Evaluation Notes\n2. Submit",
     "Error: 'Min 20 characters required.'",
     "Notes: 'Too short'"),
    ("TC-55","Case Evaluation","Valid evaluation submission",
     "A case is UNDER_EVALUATION",
     "1. Enter valid amount\n2. Select risk level\n3. Enter 20+ character notes\n4. Select recommendation\n5. Click Submit",
     "Toast: 'Evaluation submitted. Case ready for approval.' Status changes to EVALUATED.",
     "Amount: 50000\nRisk: Medium\nNotes: 'All documents verified and confirmed valid.'"),
    ("TC-56","Case Evaluation","Claim Type label in modal",
     "Any case in evaluation",
     "1. Click View\n2. Check Claim Type field",
     "Shows 'Funeral' or 'Death', not raw code.",
     "—"),

    # ── MODULE 6: Case Approval ───────────────────────────────────────────────
    ("MOD", "MODULE 6: Case Approval", "", "", "", "", ""),
    ("TC-57","Case Approval","Only evaluated cases actionable",
     "Logged in as Case Manager",
     "1. Open Case Approval",
     "Only EVALUATED cases show Approve/Deny buttons.",
     "—"),
    ("TC-58","Case Approval","Approve without notes",
     "An EVALUATED case exists",
     "1. Click Approve\n2. Leave notes blank\n3. Confirm",
     "Error: 'Required.' Case not approved.",
     "Notes: blank"),
    ("TC-59","Case Approval","Approve amount exceeds evaluated",
     "An EVALUATED case with amount set",
     "1. Click Approve\n2. Enter amount higher than evaluated amount\n3. Enter notes\n4. Confirm",
     "Server error: 'Approved amount cannot exceed evaluated amount.'",
     "If evaluated = 50000, enter 60000"),
    ("TC-60","Case Approval","Valid approval",
     "An EVALUATED case exists",
     "1. Click Approve\n2. Enter valid amount (≤ evaluated)\n3. Enter notes\n4. Click Confirm Approval",
     "Toast: 'Case approved. Payment authorised.' Status changes to APPROVED.",
     "Amount: 50000\nNotes: 'Claim verified and approved.'"),
    ("TC-61","Case Approval","Deny without notes",
     "An EVALUATED case exists",
     "1. Click Deny\n2. Leave notes blank\n3. Confirm",
     "Error: 'Required.' Case not denied.",
     "Notes: blank"),
    ("TC-62","Case Approval","Valid denial",
     "An EVALUATED case exists",
     "1. Click Deny\n2. Enter denial reason\n3. Click Confirm Denial",
     "Toast: 'Case denied.' Status changes to DENIED.",
     "Notes: 'Policy exclusion applies.'"),
    ("TC-63","Case Approval","Claim Type label in table",
     "Cases visible in approval screen",
     "1. Open Case Approval\n2. Check Type column",
     "Shows readable label, not BCP code.",
     "—"),

    # ── MODULE 7: Role-Based Access ───────────────────────────────────────────
    ("MOD", "MODULE 7: Role-Based Access Control", "", "", "", "", ""),
    ("TC-64","RBAC","Case Manager cannot access Case Registration",
     "Logged in as Case Manager",
     "1. Navigate to /case-registration directly in URL",
     "Redirected to login or access denied. 'Case Registration' not in sidebar.",
     "URL: http://localhost:9090/case-registration"),
    ("TC-65","RBAC","Claims Officer cannot access Case Approval",
     "Logged in as Claims Officer",
     "1. Check sidebar\n2. Navigate to /case-approval directly",
     "'Case Approval' link not in sidebar.",
     "URL: http://localhost:9090/case-approval"),
    ("TC-66","RBAC","Sidebar badges show pending count",
     "SUBMITTED cases exist",
     "1. Login as Claims Officer\n2. Check Case Acceptance nav link",
     "Badge shows number of SUBMITTED cases waiting for acceptance.",
     "—"),
]

# ── Write rows ────────────────────────────────────────────────────────────────
row = 2
for i, t in enumerate(TESTS):
    tc, mod, desc, pre, steps, expected, data = t

    if tc == "MOD":
        # Module separator row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = ws.cell(row=row, column=1, value=mod)
        cell.fill      = MOD_FILL
        cell.font      = MOD_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cell.border    = BORDER
        ws.row_dimensions[row].height = 22
        row += 1
        continue

    fill = ALT_FILL if i % 2 == 0 else WHITE_FILL
    values = [tc, mod, desc, pre, steps, expected, data, "", ""]
    for ci, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=ci, value=val)
        cell.fill      = fill
        cell.font      = BOLD_FONT if ci in (1, 2) else BODY_FONT
        cell.border    = BORDER
        cell.alignment = WRAP if ci not in (1, 8) else CENTER

    # Status dropdown hint
    ws.cell(row=row, column=8).value = ""   # left blank for tester
    ws.row_dimensions[row].height = 80
    row += 1

# ─────────────────────────────────────────────────────────────────────────────
# Sheet 2 – Summary
# ─────────────────────────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Summary")

summary_data = [
    ("ClaimsPro — Test Execution Summary", None),
    ("", None),
    ("Module", "# Test Cases"),
    ("Login & Authentication", 7),
    ("Dashboard", 6),
    ("Case Registration – Validations", 5),
    ("Case Registration – Dropdowns", 5),
    ("Case Registration – Policy Lookup", 5),
    ("Case Registration – Documents", 4),
    ("Case Registration – THITO Integration", 8),
    ("Case Acceptance", 10),
    ("Case Evaluation", 6),
    ("Case Approval", 7),
    ("Role-Based Access Control", 3),
    ("", None),
    ("TOTAL", 66),
]

for ri, (a, b) in enumerate(summary_data, 1):
    ca = ws2.cell(row=ri, column=1, value=a)
    if b is not None:
        cb = ws2.cell(row=ri, column=2, value=b)

    if ri == 1:
        ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        ca.font = Font(name="Calibri", bold=True, size=14, color="0F3460")
        ca.alignment = CENTER
    elif ri == 3:
        ca.fill = HDR_FILL; ca.font = HDR_FONT; ca.border = BORDER; ca.alignment = CENTER
        cb.fill = HDR_FILL; cb.font = HDR_FONT; cb.border = BORDER; cb.alignment = CENTER
    elif a == "TOTAL":
        ca.font = Font(name="Calibri", bold=True, size=11)
        ca.fill = MOD_FILL; ca.font = MOD_FONT; ca.border = BORDER; ca.alignment = CENTER
        cb.fill = MOD_FILL; cb.font = MOD_FONT; cb.border = BORDER; cb.alignment = CENTER
    elif a and b is not None:
        ca.border = BORDER; ca.font = BODY_FONT
        cb.border = BORDER; cb.font = BODY_FONT; cb.alignment = CENTER
        if ri % 2 == 0:
            ca.fill = ALT_FILL; cb.fill = ALT_FILL

ws2.column_dimensions["A"].width = 40
ws2.column_dimensions["B"].width = 16

# ─────────────────────────────────────────────────────────────────────────────
# Sheet 3 – Test Data Reference
# ─────────────────────────────────────────────────────────────────────────────
ws3 = wb.create_sheet("Test Data Reference")

ref_data = [
    ("Category", "Value", "Notes"),
    ("Claims Officer Login", "officer@claims.com / officer123", "Has access to Registration, Acceptance, Evaluation"),
    ("Case Manager Login",   "manager@claims.com / manager123", "Has access to Approval only"),
    ("Admin Login",          "admin@claims.com / admin123",     "Full access"),
    ("Application URL",      "http://localhost:9090",           "Default port"),
    ("", "", ""),
    ("ID Type – Omang",     "Code: 34",  "Sent to THITO API"),
    ("ID Type – Others",    "Code: 9",   "Sent to THITO API"),
    ("ID Type – Passport",  "Code: P",   "Sent to THITO API"),
    ("", "", ""),
    ("Gender – Male",       "Code: M",   "Sent to THITO API"),
    ("Gender – Female",     "Code: F",   "Sent to THITO API"),
    ("", "", ""),
    ("Claim Type – Funeral","Code: 15",  "THITO case starts with FUN"),
    ("Claim Type – Death",  "Code: 2",   "THITO case starts with DEA"),
    ("", "", ""),
    ("Claim Nature – Illness (Non-accident)", "Code: 1", ""),
    ("Claim Nature – Accident",              "Code: 2", ""),
    ("Claim Nature – Suicide",               "Code: 3", ""),
    ("", "", ""),
    ("Branch Code",         "101 (hardcoded)", "Always sent to THITO"),
    ("Max Document Size",   "10 MB",           "Per file"),
    ("Allowed Doc Formats", "PDF, JPG, PNG, DOCX", ""),
    ("Session Timeout",     "30 minutes",      "Configurable in application.properties"),
]

for ri, row_vals in enumerate(ref_data, 1):
    for ci, val in enumerate(row_vals, 1):
        cell = ws3.cell(row=ri, column=ci, value=val)
        cell.font   = HDR_FONT if ri == 1 else BODY_FONT
        cell.fill   = HDR_FILL if ri == 1 else (ALT_FILL if ri % 2 == 0 else WHITE_FILL)
        cell.border = BORDER
        cell.alignment = WRAP

ws3.column_dimensions["A"].width = 35
ws3.column_dimensions["B"].width = 30
ws3.column_dimensions["C"].width = 38

OUT = r"c:\ClaudeCodeAI\Claims\ClaimsPro_Test_Cases.xlsx"
wb.save(OUT)
print("Saved:", OUT)
